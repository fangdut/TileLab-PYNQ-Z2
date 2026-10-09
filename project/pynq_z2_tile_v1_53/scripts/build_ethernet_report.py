"""Build the Ethernet acceptance report from preserved real-board evidence."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1];T=R/'reports/ethernet_20260926'
def read(name):return json.loads((T/name).read_text(encoding='utf-8'))
B=read('baseline.json');F=read('scheduled.json');Q=read('board_scheduled/queue_capacity_counts.json');S=read('board_scheduled/conveyor_observation.json');C=read('recovery.json')
S1=read('board/conveyor_observation.json');S2=read('board_steady/conveyor_observation.json')
static=read('board_scheduled/static_results.json')
D=Document();sec=D.sections[0];sec.page_width=Cm(21.59);sec.page_height=Cm(27.94)
sec.top_margin=sec.bottom_margin=Cm(1.8);sec.left_margin=sec.right_margin=Cm(2)
for name,size in [('Normal',11),('Title',22),('Heading 1',17),('Heading 2',12)]:
 s=D.styles[name];s.font.name='Times New Roman' if name=='Normal' else 'Microsoft YaHei';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
 s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name=='Normal' else '微软雅黑')
D.styles['Normal'].paragraph_format.line_spacing=1.15;D.styles['Normal'].paragraph_format.space_after=Pt(6)
for s in D.styles:
 for b in list(s.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
log=[]
def p(t):log.append(t);return D.add_paragraph(t)
def h(t):log.append('\n## '+t);D.add_heading(t,2)
def page(t):D.add_page_break();log.append('\n# '+t);D.add_heading(t,1)
def table(headers,rows,widths):
 t=D.add_table(rows=1,cols=len(headers));t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 for i,row in enumerate(t.rows):
  cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
  for j,c in enumerate(row.cells):
   c.width=Cm(widths[j]);pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for side in ('top','left','bottom','right','insideH','insideV'):
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   pr.append(b);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9F0F5' if i==0 else 'FFFFFF');pr.append(shade)
   mar=OxmlElement('w:tcMar')
   for side in ('top','left','bottom','right'):
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');mar.append(e)
   pr.append(mar)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(3);para.paragraph_format.line_spacing=1.08
    for run in para.runs:run.font.size=Pt(10);run.bold=i==0
 log.append(' | '.join(headers));log.extend(' | '.join(map(str,row)) for row in rows);p('')
def eq(t):
 log.append(t);para=D.add_paragraph();math=OxmlElement('m:oMath');run=OxmlElement('m:r');text=OxmlElement('m:t');text.text=t;run.append(text);math.append(run);para._p.append(math)
footer=sec.footer.paragraphs[0];footer.alignment=2;footer.add_run('Tile V1.3  以太网评审   ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
def med(j,k):return j['summary'][k]['median']
rate=med(F,'goodput_MB_s')/med(B,'goodput_MB_s')

D.add_paragraph('瓷砖检测项目以太网测试与改进评审报告','Title')
p('PYNQ Z2   Tile V1.3 网络修订   2026 年 9 月 26 日')
p(f'本轮已完成真实开发板网线传图测试与优化。完整取帧有效速度中位数从 {med(B,"goodput_MB_s"):.2f} MB/s 提高到 {med(F,"goodput_MB_s"):.2f} MB/s，约提升 {rate:.1f} 倍。最终连续传输 303 张图像，全部通过原图、掩膜和框数据一致性校验；正式统计排除前 3 张预热帧。')
p('当前产品是瓷砖暗裂纹候选检测与人工复核原型：电脑通过 HDMI 模拟摄像头，FPGA 计算缺陷候选及红框，ARM 管理帧缓存和网络，电脑保存证据并人工复核。网络回传已从预留功能变为实板验证功能，但仍不能据此宣称达到真实工厂的漏检率要求。')
table(['实测项目','原网络基准','最终修订'],[
 ['正式样本数','24 张','300 张'],
 ['完整取帧有效速度中位数',f'{med(B,"goodput_MB_s"):.2f} MB/s',f'{med(F,"goodput_MB_s"):.2f} MB/s'],
 ['取帧耗时中位数',f'{med(B,"request_ms"):.2f} ms',f'{med(F,"request_ms"):.2f} ms'],
 ['取帧耗时 P95',f'{B["summary"]["request_ms"]["p95"]:.2f} ms',f'{F["summary"]["request_ms"]["p95"]:.2f} ms'],
 ['接收及校验编码确认周期中位数',f'{med(B,"cycle_ms"):.2f} ms',f'{med(F,"cycle_ms"):.2f} ms'],
 ['该轮新增采集帧间隙',str(B['board_deltas']['gaps']),str(F['board_deltas']['gaps'])],
 ['该轮无效帧和队列丢失','0 和 0','0 和 0']], [7.5,4.65,4.65])
p('测试环境为电脑网口直接连接 PYNQ-Z2，双方协商 1 Gbps，板地址 169.254.200.10，TCP 端口 9133，固定 1280×720、60 Hz。测速使用两个裂纹的静态合成场景、相同 BIT 和相同原始帧格式；两轮时长不同，因此帧间隙数量只说明各自观察窗口。')
p('下一步建议先解决短事件验收与按产品编号留证，再建立真实瓷砖数据集。继续堆叠演示功能的收益已经小于提高检测可信度和运行可靠性的收益。')
h('当前方案的优势')
p('FPGA 按像素流水处理，算法和区域框来源可解释；同帧图像、掩膜、参数与复核结论可以追溯；电脑断开时板端仍可暂存事件。现在已有从 HDMI 输入到 FPGA 判断、网口留图和人工复核的完整原型，适合比赛展示计算分工与工程验证。')

page('速度数字的含义和优化取舍')
h('链路速率不等于图像速度')
p('1 Gbps 表示网线链路每秒传输十亿比特，除以 8 后理论上限为 125 MB/s。协议头、确认包、等待新帧和软件处理都会消耗时间。本报告 MB 使用十进制，1 MB 等于 1,000,000 字节。')
eq('每帧字节数 = 1280 × 720 × 4 = 3,686,400')
eq('完整取帧有效速度 = 每帧字节数 ÷ 从发送请求到收完该帧的时间')
p(f'最终载荷接收阶段的中位数为 {med(F,"payload_MB_s"):.2f} MB/s；包含等待新帧的完整取帧速度为 {med(F,"goodput_MB_s"):.2f} MB/s，后者更接近实际使用。两者都不包含后续 PNG 编码和数据库落盘。P95 表示排序后约 95% 的样本耗时不高于该数值，不是最坏时间保证。')
h('为什么没有选峰值最高的版本')
p(f'第一轮优化让完整取帧速度达到约 78.72 MB/s，但出现 27 个采集帧间隙。随后限制每次发送为最多两个 TCP MSS，一帧完成后立即重新启动 DMA，并调整维护时机。最终测速窗口仍有 {F["board_deltas"]["gaps"]} 个间隙，说明轮询架构的严格连续采集还未完全解决；传输内容完整性与全时段连续采集必须分别评价。')
h('本次真正改动的计算分工')
p('裂纹候选、区域分组、面积和坐标仍全部由 FPGA 计算。本次没有增加 ARM 或电脑端实时缺陷检测。ARM 的变化是缓存配置、TCP 发送和调度；电脑的变化是数据接收、无损 PNG 编码与证据一致性检查。')
p(f'电脑端校验及 PNG 编码中位耗时从约 {med(B,"decode_png_ms"):.1f} ms 降到 {med(F,"decode_png_ms"):.1f} ms。先把反向 RGB 通道切片变为连续内存，再把 PNG 压缩级别改为 3，像素不变。单张实板原图的小型对比中，文件由 208,514 字节增至 226,134 字节，换取更短编码时间；其他画面的大小比例可能不同。')
h('必须保留的带宽边界')
eq('原始帧连续 60 fps 的带宽 = 3,686,400 × 60 = 221.184 MB/s')
p('因此当前 4 字节每像素格式无法通过千兆网无压缩回传全部 720p60 帧。FPGA 的 60 fps 检测、板端 DMA 采集、电脑取帧和浏览器刷新是不同指标。后续应把“低带宽实时预览”和“报警原始证据”分开，不能把电脑预览速度宣传成检测速度。')

page('代码改动与证据保护机制')
h('缓存与 DMA 的一致性')
p('缓存是 CPU 内部的高速副本；DMA 则直接读写内存，不会自动知道 CPU 副本是否最新。原程序关闭全部数据缓存，简单但慢。新版把 0x10000000 至 0x13FFFFFF 的共享区设为不可缓存，覆盖状态、描述符和 13 个帧槽，再开启普通程序与网络内存的缓存。GEM 网口驱动继续维护自身 DMA 缓冲的一致性。')
h('发送数据不再反复复制')
p('lwIP 是板端使用的轻量 TCP/IP 协议栈。发送头仍复制，像素直接引用队列中已保留的 DDR 帧。发送窗口从 8,192 字节增加到 65,535 字节，相应增加 TCP 段与 pbuf 描述资源。每次网络轮询最多提交 2,920 字节，使采集循环能更频繁执行。')
p('零复制有严格的生存期要求：TCP 尚可能重传的数据必须保持不变。本项目先等待整帧 TCP 确认，再接受电脑的应用层 ACK；电脑保存 PNG 并提交 SQLite 后才发 ACK。中途断线只丢弃连接，保留事件占用的帧槽。')
p('较长视频试验中曾观察到约每秒一个采集间隙，与 PHY 定时维护周期相符。最终修订将 TCP 定时器和 PHY 链路维护安排在一帧 DMA 启动后约 2 至 8 毫秒内，避开帧结束时的再次启动，并保留各阶段最大耗时诊断。该关联是调度改进依据，不能把全部历史间隙都断定为同一原因。')
h('断线后的恢复和状态反馈')
p('新增只读状态命令，电脑可查询队列深度、捕获帧数、无效帧、帧间隙与丢失计数。30 秒无进展的连接会释放，避免一个已消失的客户端长期占住服务。主机遇到网络异常最多自动重连一次；协议格式或身份错误不会被当作普通断线无限重试。ARP 和 IPv4 重组的定时维护也已补齐。')
h('无损重编码与旧证据兼容')
p('同一图像在不同 PNG 压缩级别下会产生不同文件字节。新版先校验已存文件的 SHA256，再在需要时比较解码后的原图和掩膜像素，并比较 FPGA 决策字段。这样既允许同一事件的无损重编码重试，也能拒绝改动掩膜或检测框的重复身份，且保留原人工复核结论。')
p('新增网线直连地址配置与中文启动入口，默认不修改电脑网卡设置。JTAG 下载频率降为 1 MHz，以减少本次不稳定 USB 连接对烧录的影响。下载仍只修改板上易失运行状态，不写 Flash 或 SD 卡。')

page('本轮真实开发板验收')
table(['场景','候选像素','区域框数','掩膜与框参考比对'],[
 [name,row['hardware']['candidate_pixels'],len(row['hardware']['boxes']),'一致' if row['passed'] else '失败']
 for name,row in zip(['正常表面','细裂纹','反光干扰','宽裂纹','两个分离裂纹','孤立小点'],static)], [5.3,3.0,3.0,5.5])
p('每种场景都经过电脑 HDMI 输入、真实 FPGA、DDR、网口和电脑接收链路。离线独立参考只用于验收，不参与实时检测。六类场景掩膜差异均为 0 像素，区域坐标与面积全部一致。')
h('事件缓存与断线试验')
p(f'0.6 秒缺陷和 0.6 秒正常画面交替，产生 {Q["event_delta"]} 次事件。暂停电脑接收时板端保留 12 张，额外 {Q["dropped_delta"]} 张明确计为队列满丢失；随后 12 张图像和缺陷框全部正确取回，最终队列清空，ACK 错误为 0。故意制造的 4 张丢失是容量边界测试结果，不是正常演示的漏传。')
p('在真实网线上中断半帧接收后，重连得到相同启动编号、事件编号和原始帧 SHA256；重复 ACK 不会释放下一事件；错误协议头不会释放证据；空闲超时后可重新连接。四项均通过。这里测试的是断开 TCP 连接，不等于已经完成物理拔网线和断电恢复验收。')
h('移动流水线观察')
p(f'最终固件观察 {S["duration_seconds"]:.1f} 秒，板端帧号增加 {S["frames_delta"]}，完整采集增加 {S["capture_delta"]}，电脑收到 {S["received_delta"]} 张，保存 {S["records_added"]} 条事件；新增无效帧 {S["invalid_delta"]}、采集帧间隙 {S["gap_delta"]}、队列满丢失 {S["dropped_delta"]}。'+('本轮观察通过。' if S['passed'] else '严格连续采集仍有未通过项，需继续排查。'))
p('每 2 秒的采样还观察到两段电脑接收计数不增加的区间，约为第 56.47 至 70.59 秒和第 157.28 至 173.48 秒，分别至少 14.12 秒和 16.20 秒，随后恢复。最终错误字段为空不代表全程无停顿；目前未记录足够的超时、重连和 TCP 重传日志来确定原因，预览实时性尚未验收通过。')
p(f'保留失败记录：中间固件的 {S1["duration_seconds"]:.1f} 秒观察有 {S1["gap_delta"]} 个间隙；4 秒预热后延长观察 {S2["duration_seconds"]:.1f} 秒仍有 {S2["gap_delta"]} 个间隙，因此没有把问题简单归为播放器启动。最终结果不能覆盖或删除这些失败证据。')
p('电脑端 22 项测试通过，覆盖协议拆包、截断、重连、保存后确认、编号去重、掩膜篡改及 PNG 重编码。RTL、BIT 与检测参数保持原有已验证版本，本轮没有重新声称执行全套 RTL 仿真。')
p('这些短时间实测不能替代数小时或数天的连续运行验收。旧版 100 毫秒视频短事件曾有成功与失败轮次，本轮尚未将它重新定义为已稳定通过的能力。')

page('项目审核发现的主要不足')
h('最高优先级是检测可信度')
p('目前全部通过数据仍是合成素材，无法给出工厂准确率。四方向暗谷算法适合暗裂纹，却可能把纹理接缝、污渍和阴影当候选，也可能漏掉低对比度、特别宽或亮色缺陷。17 像素窗口和阈值 36 对相机距离、光照和表面纹理敏感。下一步应先收集真实好品与裂纹、反光、污渍、花纹照片，分开调参集和独立测试集。')
h('事件尚未等同于每一块产品')
p('当前按报警由无到有留图，连续 3 张正常采集帧后重新触发。连续坏品之间若没有正常间隔，可能合成一个事件。实看本轮帧 26310 的自动截图时，裂纹刚进入 ROI，只保存了 43 个候选像素、坐标 [80,393,107,408] 的小框，缺陷尚未完整进入画面。应增加产品编号、入站与离站状态，在 FPGA 中按缺陷面积和完整入区程度选择代表帧；不能只保留第一次报警的画面。')
h('短事件和复位跨时钟域仍需闭环')
p('历史 100 毫秒测试不一致的问题尚未完全隔离。建议给 HDMI 测试视频每帧加入可读编号，联合记录输入帧号、PL 报警、DMA 完成和事件编号，判断遗漏发生在哪一段。最终实现报告仍有 5 个跨时钟域不安全端点，集中于 HDMI 相关域；时序通过不代表跨域电路已经安全，需逐项定位并验证异步置位、同步释放及握手机制。')
p('本轮诊断曾记录单次网络服务最长约 3.295 毫秒，已超过 720p60 约 0.667 毫秒的垂直消隐时间。这支持“ARM 轮询存在帧边界调度风险”的判断，但尚未把每个间隙与具体调用逐一对应。下一轮应优先用预置 DMA 缓冲和硬件帧边界控制解除这个依赖。')
h('检测框与显示方式仍有边界')
p('32×32 网格占用连通会合并相邻细裂纹；最多 8 个区域，超过数量会给出溢出标记。HDMI 叠框使用上一完整帧结果，移动时可能有位置偏差；保存证据的原图与框已经按同一帧对应。下一步应按需要缩小网格或改用行程连通域，并用帧缓存或运动补偿解决显示对齐。')
h('缓存和长期运行')
p('12 张 DDR 缓存不保断电，电脑存档达到 500 条会停止继续采集保存，当前没有完整的自动归档和恢复流程。只靠更快网络不能覆盖电脑长时间离线。还需要存储配额、磁盘满演练、持久化配置、看门狗与可启动镜像。网络按可信直连实验环境设计，不含认证与加密。')
p('现有实现约占 38.64% LUT、32.30% FF、41.07% BRAM，最差建立时间余量只有 0.087 ns。增加多尺度和缓存控制前要重新估算资源、流水级与关键路径。电脑端大量历史缩略图也应分页加载，避免浏览复核记录影响预览响应。')

page('建议的下一轮实施方案')
h('先做能够量化验收的产品流程')
p('建议先改造帧采集衔接，再实现产品编号与双通道取图。采集采用硬件自动衔接或 DMA 描述符队列，提前准备下一缓冲，避免 ARM 必须恰好在帧边界处理完网络。随后 FPGA 按产品窗口分配编号，离站时汇总结论；电脑分别接收实时缩小预览和完整报警证据，均关联同一产品号。')
table(['优先级','具体工作','建议验收方式'],[
 ['P0','自动衔接 DMA 与短事件定位','固定序列重复 10 轮，逐帧编号与事件清单核对'],
 ['P0','定位电脑取图长停顿','记录重连和重传，区分请求等待与载荷停顿，复测长视频'],
 ['P0','真实图像与独立测试集','先做约 200 张试验集，记录来源、批次、光照和标签'],
 ['P1','产品编号与最佳证据帧','连续坏品有独立记录，保存缺陷完整进入检测区的帧'],
 ['P1','预览与完整证据分离','积压 12 条事件时仍能看见有时戳的当前预览'],
 ['P1','跨域复位和失锁恢复','逐项关闭 CDC 风险，测试 HDMI 拔插与输入重启'],
 ['P2','多尺度裂纹和形态筛选','与固定算法对照，报告漏检、误报、资源和时序'],
 ['P2','归档和掉电恢复','磁盘满、程序崩溃、重启及恢复日志可追溯']], [1.8,5.8,9.2])
p('约 200 张是启动采样建议，不足以支撑工业可靠性承诺。不要把同一视频的相邻帧随机分到训练与测试两边；应按瓷砖或生产批次隔离，防止测试集过于相似而高估效果。')
h('在 FPGA 中继续优化算法')
p('可以并行计算多个半径的暗谷响应，并按连通区域的长度、宽度、面积、填充率等筛选；背景照明缓慢变化时可增加局部归一化。但这些改动必须用真实反例验证。引入 YOLO 并不是必经步骤，当前 Zynq-7020 的资源、带宽和比赛展示需求更适合先把可解释的流式算法与完整产品流程做扎实。')
p('建议分别统计真实缺陷中检出的比例即召回率，以及所有报警中确为缺陷的比例即精确率；同时报告每千块好品误报数和每块产品留证完整率。区域框的 IoU 可以辅助衡量定位，但细长裂纹应结合像素召回，避免只看大框造成误解。')

page('连接操作和复现入口')
h('保持这四条连接')
p('电脑 HDMI 输出接板 HDMI IN，板 HDMI OUT 接显示器，PROG UART 用可传数据的 Micro-USB 线接电脑，板上 RJ45 网口用普通网线直接接电脑网口。可靠供电后，Windows 扩展 HDMI 输出固定为 1280×720、60 Hz。本次串口为 COM20，实际使用时编号可能变化。')
p('电脑有线网卡本次自动地址为 169.254.173.242，掩码 /16；板端默认为 169.254.200.10，掩码 /16，不需要路由器或互联网，也不需要关闭 Wi-Fi。换电脑后先检查其有线 IPv4 地址是否在相同网段。')
h('烧录和启动')
p('工程根目录：D:/Vivado/pynq_z2_tile_v1_5。先运行同名 .sdk 文件夹中的 program_ethernet.bat，等待“Programmed”成功信息；再双击根目录的“启动以太网工作站.bat”。浏览器地址是 http://127.0.0.1:8770。不能让 USB 工作站和网口工作站同时占用同一服务端口。')
p('更换为路由器 DHCP 模式时，可在命令行给 program_ethernet.bat 追加 dhcp，再从 115200 波特率串口读取地址，并把地址传给启动批处理。DHCP 模式本轮没有物理路由器验收。默认直连地址和网络选择不会改动电脑已有网卡配置。')
h('复现实测')
p('关闭或暂停工作站并释放其网络连接后，运行 scripts/benchmark_ethernet.py --label repeat --frames 300。基准工具使用真实板端帧，检验内容后才 ACK；报警事件会先存档。结果写入 reports/ethernet_20260926。不要同时运行测速工具和工作站采集，两者使用同一板端 TCP 服务。')
p('六类画面和容量测试入口为 scripts/board_v13_acceptance.py，容量模式加 --queue；连续视频观察为 scripts/board_soak.py；断线恢复为 scripts/check_ethernet_recovery.py。测试时画面切换属于输入刺激，完成后可选择普通流水线视频演示。')
h('保存和回退')
p('修改前备份为 D:/Vivado/_baseline_backups/tile_v13_before_ethernet_20260926.zip，大小 215,743,085 字节。原始 HDL、BIT、旧软件、数据与复核记录均在其中。恢复时解压到另一目录进行核对，避免覆盖本次新增复核记录。原基准网口 ELF 另存于 reports/ethernet_20260926/baseline。')
p('关键证据：baseline.json、scheduled.json、recovery.json、board_scheduled/static_results.json、board_scheduled/queue_capacity_counts.json、board_scheduled/conveyor_observation.json，以及 host_tests.log。哈希和烧录记录见 ETHERNET_REVIEW_MANIFEST.json；旧版清单保留历史验收含义。')
p('技术依据：项目 network.c、event_queue.h、valley_reference.py 与 Vivado 报告；AMD 随工具安装的 lwIP 2.2.0 GEM 驱动及 Cortex-A9 缓存实现；lwIP Raw API 官方说明 https://www.nongnu.org/lwip/2_1_x/group__tcp__raw.html；AMD lwIP 说明 https://xilinx-wiki.atlassian.net/wiki/spaces/A/pages/18842366/Standalone+LWIP+library。')
p('DMA 队列方案依据 AMD PG021 的 Scatter/Gather 描述符机制：https://docs.amd.com/r/en-US/pg021_axi_dma/Scatter-Gather-Descriptor。采用该方案还需改造当前帧捕获控制、描述符所有权和事件缓冲保护，不能只切换一个 IP 参数。')

target=R/'docs/Tile_V1.3_以太网实测与项目改进评审报告.docx'
D.save(target);target.with_suffix('.md').write_text('\n\n'.join(log),encoding='utf-8')
print(target)
